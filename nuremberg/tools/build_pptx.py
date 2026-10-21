"""Сборка PowerPoint: снимок слайда на фоне + видео в колонке справа + прозрачный оверлей «рилс» поверх.

python3 build_pptx.py <папка со снимками из shots.mjs> <выходной .pptx>
Видео берутся из ../media/<clip>.mp4; у каждого слайда своя точка старта, видео идёт само, без звука и по кругу.
"""
import json
import re
import subprocess
import sys
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.util import Emu

HERE = Path(__file__).resolve().parent
ROOT = HERE.parent
shots, out = Path(sys.argv[1]), Path(sys.argv[2])

html = (ROOT / 'index.html').read_text(encoding='utf-8')
clips = re.findall(r'<section class="slide[^"]*"[^>]*data-clip="(\w+)"', html)
speech = (ROOT / 'текст-доклада.md').read_text(encoding='utf-8')
notes = [m.strip() for m in re.findall(r'\*\*Слайд \d+\.[^*]*\*\*\n(.*?)(?=\n\*\*Слайд|\n---|\Z)', speech, re.S)]
assert len(notes) == len(clips), (len(notes), len(clips))
reel = json.loads((shots / 'reel.json').read_text())

W, H = 12192000, 6858000
K = W / 1600                      # EMU на пиксель сцены 1600×900
B = 3                             # рамка колонки
vx, vy, vw, vh = (round((reel['x'] + B) * K), round((reel['y'] + B) * K),
                  round((reel['w'] - 2 * B) * K), round((reel['h'] - 2 * B) * K))


def duration(clip):
    r = subprocess.run(['ffprobe', '-v', 'error', '-show_entries', 'format=duration', '-of', 'csv=p=0',
                        str(ROOT / 'media' / f'{clip}.mp4')], capture_output=True, text=True, check=True)
    return float(r.stdout)


def poster(clip, t):
    p = shots / f'poster_{clip}_{t:.1f}.jpg'
    if not p.exists():
        subprocess.run(['ffmpeg', '-loglevel', 'error', '-y', '-ss', f'{t:.2f}', '-i', str(ROOT / 'media' / f'{clip}.mp4'),
                        '-frames:v', '1', '-q:v', '3', str(p)], check=True)
    return p


TIMING = '''<p:timing xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main"><p:tnLst><p:par><p:cTn id="1" dur="indefinite" restart="never" nodeType="tmRoot"><p:childTnLst>
<p:par><p:cTn id="2" fill="hold"><p:stCondLst><p:cond delay="indefinite"/><p:cond evt="onBegin" delay="0"><p:tn val="2"/></p:cond></p:stCondLst><p:childTnLst>
<p:par><p:cTn id="3" fill="hold"><p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst>
<p:par><p:cTn id="4" presetID="1" presetClass="mediacall" presetSubtype="0" fill="hold" nodeType="afterEffect"><p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst>
<p:cmd type="call" cmd="playFrom({start})"><p:cBhvr><p:cTn id="5" dur="{dur}" fill="hold"/><p:tgtEl><p:spTgt spid="{spid}"/></p:tgtEl></p:cBhvr></p:cmd>
</p:childTnLst></p:cTn></p:par></p:childTnLst></p:cTn></p:par></p:childTnLst></p:cTn></p:par>
<p:video><p:cMediaNode vol="0" mute="1"><p:cTn id="6" repeatCount="indefinite" fill="hold" display="0"><p:stCondLst><p:cond delay="indefinite"/></p:stCondLst></p:cTn><p:tgtEl><p:spTgt spid="{spid}"/></p:tgtEl></p:cMediaNode></p:video>
</p:childTnLst></p:cTn></p:par></p:tnLst></p:timing>'''

prs = Presentation()
prs.slide_width, prs.slide_height = Emu(W), Emu(H)
blank = prs.slide_layouts[6]
durs = {c: duration(c) for c in set(clips)}
seen = {}
for i, clip in enumerate(clips, 1):
    s = prs.slides.add_slide(blank)
    bg = shots / f'slide_{i:02d}.jpg'                   # фон в JPEG: файл заметно легче, чем с PNG
    subprocess.run(['ffmpeg', '-loglevel', 'error', '-y', '-i', str(shots / f'slide_{i:02d}.png'), '-q:v', '2', str(bg)], check=True)
    s.shapes.add_picture(str(bg), 0, 0, W, H)
    k = seen.get(clip, 0); seen[clip] = k + 1
    start = (k * 7.3) % (durs[clip] - 2)               # каждый слайд — с нового места ролика
    mv = s.shapes.add_movie(str(ROOT / 'media' / f'{clip}.mp4'), vx, vy, vw, vh,
                            poster_frame_image=str(poster(clip, start)), mime_type='video/mp4')
    mv.name = f'{clip}.mp4'
    geom = mv._element.spPr.find('{http://schemas.openxmlformats.org/drawingml/2006/main}prstGeom')
    geom.set('prst', 'roundRect')
    av = geom.find('{http://schemas.openxmlformats.org/drawingml/2006/main}avLst')
    etree.SubElement(av, '{http://schemas.openxmlformats.org/drawingml/2006/main}gd', name='adj', fmla='val 6000')
    s.shapes.add_picture(str(shots / f'over_{i:02d}.png'), round(reel['x'] * K), round(reel['y'] * K),
                         round(reel['w'] * K), round(reel['h'] * K))
    # заменяем «запуск по щелчку» на автозапуск без звука по кругу
    sld = s._element
    for t in sld.findall('{http://schemas.openxmlformats.org/presentationml/2006/main}timing'):
        sld.remove(t)
    sld.append(etree.fromstring(TIMING.format(start=f'{start:.1f}', dur=int(durs[clip] * 1000), spid=mv.shape_id)))
    s.notes_slide.notes_text_frame.text = notes[i - 1]
prs.save(out)
print('saved', out, len(clips), 'slides')
